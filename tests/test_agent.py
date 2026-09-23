from __future__ import annotations

import csv
import json
import tempfile
import unittest
from email import policy
from email.parser import BytesParser
from pathlib import Path
from unittest.mock import patch

from docx import Document

from internship_agent.config import load_config
from internship_agent.extract import extract_job
from internship_agent.pipeline import run_pipeline
from internship_agent.privacy import scan_repository
from internship_agent.resume import reorder_resume_copy
from internship_agent.sources import IngestedSource
from internship_agent.models import SourceDefinition


FIXTURE_DIR = Path(__file__).parent / "fixtures"


def _minimal_pdf(path: Path) -> Path:
    path.write_bytes(
        b"%PDF-1.1\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Count 0/Kids[]>>endobj\n"
        b"trailer<</Root 1 0 R>>\n%%EOF\n"
    )
    return path


class ExtractionTests(unittest.TestCase):
    def test_extracts_required_fields(self) -> None:
        text = (FIXTURE_DIR / "job-post.txt").read_text(encoding="utf-8")
        source = SourceDefinition("fixture", "text", text, "测试来源")
        ingested = IngestedSource(
            definition=source,
            text=text,
            raw_path=Path("raw.txt"),
            extracted_path=Path("extracted.txt"),
            source_hash="abc123",
        )
        job = extract_job(ingested)
        self.assertEqual(job.company, "示例资本管理有限公司")
        self.assertEqual(job.title, "行业研究实习生")
        self.assertEqual(job.location, "上海")
        self.assertEqual(job.email, "recruit@example.com")
        self.assertEqual(job.deadline, "2027-01-31")
        self.assertIn("硕士在读", job.requirements)
        self.assertEqual(job.email_subject_rule, "姓名-学校-行业研究实习")
        self.assertEqual(job.attachment_naming_rule, "姓名-学校-岗位.pdf")


class ResumeTests(unittest.TestCase):
    def test_reorders_only_configured_evidence_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            base = root / "base.docx"
            output = root / "output.docx"
            document = Document()
            document.add_paragraph("示例经历")
            document.add_paragraph("• 政策研究：整理政策资料")
            document.add_paragraph("• 行业研究：撰写行业报告")
            document.add_paragraph("其他内容")
            document.save(base)

            class Group:
                paragraph_indices = (1, 2)
                evidence_ids = ("E01", "E02")

            before = base.read_bytes()
            order = reorder_resume_copy(
                base,
                output,
                (Group(),),
                "岗位要求开展行业研究",
                ("行业研究", "政策研究"),
            )
            result = Document(output)
            self.assertEqual(result.paragraphs[1].text, "• 行业研究：撰写行业报告")
            self.assertEqual(result.paragraphs[2].text, "• 政策研究：整理政策资料")
            self.assertEqual(result.paragraphs[3].text, "其他内容")
            self.assertEqual(order, ("E02", "E01"))
            self.assertEqual(base.read_bytes(), before)


class PipelineTests(unittest.TestCase):
    def _write_config(self, root: Path) -> Path:
        case = root / "case"
        case.mkdir()
        source_file = root / "job-post.txt"
        source_file.write_text((FIXTURE_DIR / "job-post.txt").read_text(encoding="utf-8"), encoding="utf-8")
        sources = root / "sources.local.toml"
        sources.write_text(
            f'''[[sources]]\nid = "fixture"\ntype = "file"\nname = "测试来源"\nvalue = "{source_file.as_posix()}"\nenabled = true\n''',
            encoding="utf-8",
        )

        base = case / "base.docx"
        document = Document()
        document.add_paragraph("示例简历")
        document.add_paragraph("• 政策研究：整理政策资料")
        document.add_paragraph("• 行业研究：撰写行业报告")
        document.save(base)

        config = root / "config.local.toml"
        config.write_text(
            f'''
[case]
case_dir = "{case.as_posix()}"
sources_file = "{sources.as_posix()}"

[identity]
name = "张同学"
school = "示例大学"
major = "金融学"
grade = "研一"
availability = "两周内"

[preferences]
cities = ["上海"]
target_roles = ["行业研究实习"]
available_days_per_week = 5
internship_months = 6

[matching]
preparation_threshold = 70

[resume]
soffice_command = "soffice"

[[roles]]
name = "行业研究实习"
keywords = ["行业研究", "政策研究", "数据分析"]
evidence_ids = ["E01", "E02"]
base_docx = "base.docx"

[[roles.groups]]
paragraph_indices = [1, 2]
evidence_ids = ["E01", "E02"]
''',
            encoding="utf-8",
        )
        return config

    def test_pipeline_is_idempotent_and_builds_review_gated_package(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            config_path = self._write_config(root)
            config = load_config(config_path)

            def fake_convert(docx_path: Path, output_dir: Path, command: str) -> Path:
                return _minimal_pdf(output_dir / f"{docx_path.stem}.pdf")

            with patch("internship_agent.resume.convert_to_pdf", side_effect=fake_convert), patch(
                "internship_agent.resume.verify_pdf",
                return_value={"page_count": 1, "extractable_text": True, "visual_review_required": True},
            ):
                first = run_pipeline(config)
                second = run_pipeline(config)

            self.assertEqual(first["applications_prepared"], 1)
            self.assertEqual(second["applications_prepared"], 1)
            jobs = config.jobs_file.read_text(encoding="utf-8").strip().splitlines()
            self.assertEqual(len(jobs), 1)
            with config.tracker_file.open("r", encoding="utf-8-sig", newline="") as handle:
                tracker_rows = list(csv.DictReader(handle))
            self.assertEqual(len(tracker_rows), 1)
            self.assertEqual(tracker_rows[0]["status"], "待确认")
            self.assertEqual(tracker_rows[0]["user_confirmed"], "false")

            job = json.loads(jobs[0])
            draft_dir = config.drafts_dir / job["job_id"]
            approval = json.loads((draft_dir / "approval.json").read_text(encoding="utf-8"))
            self.assertFalse(approval["external_action_performed"])
            self.assertFalse(approval["user_confirmed"])
            message = BytesParser(policy=policy.default).parsebytes((draft_dir / "email.eml").read_bytes())
            attachments = list(message.iter_attachments())
            self.assertEqual(len(attachments), 1)
            self.assertEqual(attachments[0].get_content_type(), "application/pdf")
            self.assertEqual(message["To"], "recruit@example.com")


class PrivacyTests(unittest.TestCase):
    def test_repository_has_no_sensitive_artifacts(self) -> None:
        repo = Path(__file__).resolve().parents[1]
        self.assertEqual(scan_repository(repo), [])

    def test_no_email_transport_module_is_imported(self) -> None:
        repo = Path(__file__).resolve().parents[1]
        code = "\n".join(path.read_text(encoding="utf-8") for path in (repo / "internship_agent").glob("*.py"))
        self.assertNotIn("import smtplib", code)
        self.assertNotIn("requests.post", code)


if __name__ == "__main__":
    unittest.main()
